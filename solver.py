import json
import boto3
import openpyxl
from io import BytesIO
from ortools.linear_solver import pywraplp
 
s3 = boto3.client("s3")
 
S3_BUCKET = "230-hcp-territories-data"
S3_KEY = "Input_UK_HCP_Universe_MockData_TerritoryZero_v1.xlsx"
 
 
def load_sheet_as_dicts(wb, sheet_name):
    ws = wb[sheet_name]
    headers = [cell.value for cell in ws[1]]
    return [dict(zip(headers, row)) for row in ws.iter_rows(min_row=2, values_only=True)]
 
 
def to_float(v):
    try:
        return float(v) if v is not None else 0.0
    except (TypeError, ValueError):
        return 0.0
 
 
def truthy(v):
    if v is None:
        return False
    return str(v).strip().lower() in ("true", "yes", "y", "1")
 
 
def lambda_handler(event, context):
    body = json.loads(event["body"]) if "body" in event else event
    brief = body.get("brief", body)
    run_id = body.get("run_id", "manual")
    try:
        num_territories = int(brief.get("num_territories", 8))
    except (TypeError, ValueError):
        num_territories = 8
 
    obj = s3.get_object(Bucket=S3_BUCKET, Key=S3_KEY)
    wb = openpyxl.load_workbook(BytesIO(obj["Body"].read()), data_only=True)
 
    hcp_rows = load_sheet_as_dicts(wb, "HCP_Universe")
    geo_rows = load_sheet_as_dicts(wb, "Geography_Reference")
    geo_lookup = {g.get("Postcode_Sector"): g for g in geo_rows if g.get("Postcode_Sector")}
 
    # --- Build per-HCP records, with every field we need downstream ---
    hcps = []
    for r in hcp_rows:
        status = (r.get("Record_Status") or "").strip().lower()
        if status and status != "active":
            continue
        sector = r.get("Postcode_Sector")
        hcps.append({
            "hcp_id": r.get("HCP_ID"),
            "first_name": r.get("First_Name"),
            "last_name": r.get("Last_Name"),
            "primary_specialty": r.get("Primary_Specialty"),
            "primary_hco_name": r.get("Primary_HCO_Name"),
            "best_segment": r.get("Best_Segment"),
            "postcode_sector": sector,
            "nhs_region": r.get("NHS_Region"),
            "target_flag": truthy(r.get("Target_Flag")),
            "kol_flag": truthy(r.get("KOL_Flag")),
            "is_referral_centre": truthy(r.get("Is_Referral_Centre")),
            "workload_units": to_float(r.get("Workload_Units")),
            "value_units": to_float(r.get("Value_Units")),
            "eligible_patients_est": to_float(r.get("Eligible_Patients_Est")),
        })
 
    # --- Roll up to postcode sector ---
    sector_agg = {}
    for h in hcps:
        sec = h["postcode_sector"] or "UNKNOWN"
        a = sector_agg.setdefault(sec, {
            "hcp_count": 0, "target_hcp": 0, "kol_count": 0,
            "referral_centre_hcp": 0, "workload_units": 0.0,
            "value_units": 0.0, "eligible_patients_est": 0.0,
        })
        a["hcp_count"] += 1
        a["target_hcp"] += 1 if h["target_flag"] else 0
        a["kol_count"] += 1 if h["kol_flag"] else 0
        a["referral_centre_hcp"] += 1 if h["is_referral_centre"] else 0
        a["workload_units"] += h["workload_units"]
        a["value_units"] += h["value_units"]
        a["eligible_patients_est"] += h["eligible_patients_est"]
 
    sectors = list(sector_agg.keys())
    n = len(sectors)
    total_workload = sum(a["workload_units"] for a in sector_agg.values())
    target = total_workload / num_territories if num_territories else total_workload
 
    # --- Balance sectors across territories by workload ---
    solver = pywraplp.Solver.CreateSolver("CBC")
    x = {(i, t): solver.BoolVar(f"x_{i}_{t}") for i in range(n) for t in range(num_territories)}
    for i in range(n):
        solver.Add(sum(x[i, t] for t in range(num_territories)) == 1)
 
    max_dev = solver.NumVar(0, solver.infinity(), "max_dev")
    for t in range(num_territories):
        load = sum(sector_agg[sectors[i]]["workload_units"] * x[i, t] for i in range(n))
        solver.Add(load - target <= max_dev)
        solver.Add(target - load <= max_dev)
    solver.Minimize(max_dev)
    solver.SetTimeLimit(20000)
    solver.Solve()
 
    t_ids = [f"T{t + 1}" for t in range(num_territories)]
    t_names = [f"Territory_{t + 1}" for t in range(num_territories)]
    sector_to_territory = {}
    for i in range(n):
        for t in range(num_territories):
            if x[i, t].solution_value() > 0.5:
                sector_to_territory[sectors[i]] = t
                break
 
    # --- Sheet1: sector_territory_mapping ---
    sheet1_rows = []
    for sec in sectors:
        geo = geo_lookup.get(sec, {})
        a = sector_agg[sec]
        t = sector_to_territory[sec]
        sheet1_rows.append({
            "postcode_sector": sec,
            "postcode_district": geo.get("Postcode_District"),
            "postcode_area": geo.get("Postcode_Area"),
            "post_town": geo.get("Post_Town"),
            "county_unitary_authority": geo.get("County_Unitary_Authority"),
            "country": geo.get("NHS_Nation"),
            "nhs_region": geo.get("NHS_Region"),
            "icb_or_health_board": geo.get("ICB_Or_Health_Board"),
            "latitude": to_float(geo.get("Latitude")),
            "longitude": to_float(geo.get("Longitude")),
            "urban_rural_class": geo.get("Urban_Rural_Class"),
            "hcp_count": a["hcp_count"],
            "target_hcp": a["target_hcp"],
            "workload_units": round(a["workload_units"], 1),
            "value_units": round(a["value_units"], 1),
            "eligible_patients_est": round(a["eligible_patients_est"], 1),
            "territory_id": t_ids[t],
            "territory_name": t_names[t],
        })
 
    # --- Sheet2: Territory_summary ---
    territory_summary = {}
    for t in range(num_territories):
        members = [sec for sec in sectors if sector_to_territory[sec] == t]
        nation_counts = {}
        lat_sum = lon_sum = wt_sum = 0.0
        areas, icbs = set(), set()
        agg = {"hcp_count": 0, "target_hcp": 0, "kol_count": 0,
               "referral_centre_hcp": 0, "workload_units": 0.0,
               "value_units": 0.0, "eligible_patients_est": 0.0}
        for sec in members:
            geo = geo_lookup.get(sec, {})
            a = sector_agg[sec]
            nation = geo.get("NHS_Nation")
            if nation:
                nation_counts[nation] = nation_counts.get(nation, 0) + 1
            lat, lon = to_float(geo.get("Latitude")), to_float(geo.get("Longitude"))
            w = a["workload_units"] or 1.0
            lat_sum += lat * w
            lon_sum += lon * w
            wt_sum += w
            if geo.get("Postcode_Area"):
                areas.add(geo.get("Postcode_Area"))
            if geo.get("ICB_Or_Health_Board"):
                icbs.add(geo.get("ICB_Or_Health_Board"))
            for k in agg:
                agg[k] += a[k]
 
        dominant_nation = max(nation_counts, key=nation_counts.get) if nation_counts else None
        territory_summary[t_ids[t]] = {
            "territory_id": t_ids[t],
            "territory_name": t_names[t],
            "dominant_nhs_nation": dominant_nation,
            "centroid_latitude": round(lat_sum / wt_sum, 5) if wt_sum else None,
            "centroid_longitude": round(lon_sum / wt_sum, 5) if wt_sum else None,
            "postcode_sectors": members,
            "postcode_areas": sorted(areas),
            "icbs_or_boards_touched": sorted(icbs),
            "hcp_count": agg["hcp_count"],
            "target_hcp_count": agg["target_hcp"],
            "kol_count": agg["kol_count"],
            "referral_centre_hcp": agg["referral_centre_hcp"],
            "eligible_patients_est": round(agg["eligible_patients_est"], 1),
            "workload_units": round(agg["workload_units"], 1),
            "value_units": round(agg["value_units"], 1),
        }
 
    # --- Sheet3: Hcp_territory_assignment (full HCP-level detail) ---
    assignments = []
    for h in hcps:
        t = sector_to_territory.get(h["postcode_sector"] or "UNKNOWN")
        if t is None:
            continue
        geo = geo_lookup.get(h["postcode_sector"], {})
        assignments.append({
            "hcp_id": h.get("hcp_id"),
            "first_name": h.get("first_name"),
            "last_name": h.get("last_name"),
            "primary_specialty": h.get("primary_specialty"),
            "primary_hco_name": h.get("primary_hco_name"),
            "best_segment": h.get("best_segment"),
            "postcode_sector": h.get("postcode_sector"),
            "nhs_region": h.get("nhs_region") or geo.get("NHS_Region"),
            "is_target": h.get("target_flag", False),
            "is_kol": h.get("kol_flag", False),
            "is_referral_centre": h.get("is_referral_centre", False),
            "territory_id": t_ids[t],
            "territory_name": t_names[t],
            "workload_units": h.get("workload_units"),
            "value_units": h.get("value_units"),
            "eligible_patients_est": h.get("eligible_patients_est"),
        })
    assignments.sort(key=lambda a: (a["territory_name"], str(a["hcp_id"])))
 
    result = {
        "hcp_count": len(hcps),
        "assignments": assignments,
        "territory_summary": territory_summary,
        "sector_mapping": sheet1_rows,
    }
 
    json_key = f"outputs/{run_id}/territory_alignment.json"
    xlsx_key = f"outputs/{run_id}/territory_alignment.xlsx"
 
    s3.put_object(Bucket=S3_BUCKET, Key=json_key,
                   Body=json.dumps(result), ContentType="application/json")
 
    # --- Excel: 4 sheets ---
    out = openpyxl.Workbook()
 
    ws1 = out.active
    ws1.title = "sector_territory_mapping"
    cols1 = list(sheet1_rows[0].keys()) if sheet1_rows else []
    ws1.append(cols1)
    for row in sheet1_rows:
        ws1.append([row.get(c) for c in cols1])
 
    ws2 = out.create_sheet("Territory_summary")
    cols2 = ["territory_id", "territory_name", "dominant_nhs_nation",
              "centroid_latitude", "centroid_longitude", "postcode_sectors",
              "postcode_areas", "icbs_or_boards_touched", "hcp_count",
              "target_hcp_count", "kol_count", "referral_centre_hcp",
              "eligible_patients_est", "workload_units", "value_units"]
    ws2.append(cols2)
    for s_ in territory_summary.values():
        ws2.append([
            s_["territory_id"], s_["territory_name"], s_["dominant_nhs_nation"],
            s_["centroid_latitude"], s_["centroid_longitude"],
            ", ".join(s_["postcode_sectors"]), ", ".join(s_["postcode_areas"]),
            ", ".join(s_["icbs_or_boards_touched"]), s_["hcp_count"],
            s_["target_hcp_count"], s_["kol_count"], s_["referral_centre_hcp"],
            s_["eligible_patients_est"], s_["workload_units"], s_["value_units"],
        ])
 
    ws3 = out.create_sheet("Hcp_territory_assignment")
    cols3 = list(assignments[0].keys()) if assignments else []
    ws3.append(cols3)
    for a in assignments:
        ws3.append([a.get(c) for c in cols3])
 
    ws4 = out.create_sheet("veeva_align_import")
    ws4.append(["territory_name", "geography_key", "action"])
    for row in sheet1_rows:
        ws4.append([row["territory_name"], row["postcode_sector"], "ADD"])
 
    buf = BytesIO()
    out.save(buf)
    s3.put_object(Bucket=S3_BUCKET, Key=xlsx_key, Body=buf.getvalue(),
                   ContentType="application
 
